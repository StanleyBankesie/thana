import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
import os

def create_account_template(output_paths):
    wb = Workbook()
    
    # 1. Lookup sheet for data validation dropdowns
    lookup_sheet = wb.active
    lookup_sheet.title = "LookupData"
    
    classes = ["Asset", "Liability", "Equity", "Income", "Expense"]
    currencies = ["GHS", "USD", "EUR", "GBP", "NGN", "KES", "CAD", "CNY"]
    postable_opts = ["Yes", "No"]
    balance_types = ["DEBIT", "CREDIT"]
    sample_groups = [
        "Current Assets",
        "Cash and Cash Equivalents",
        "Bank Accounts",
        "Accounts Receivable",
        "Fixed Assets",
        "Current Liabilities",
        "Accounts Payable",
        "Owner's Equity",
        "Operating Revenue",
        "Operating Expenses",
        "Administrative Expenses"
    ]

    lookup_sheet["A1"] = "ACCOUNT_CLASS"
    lookup_sheet["B1"] = "CURRENCY"
    lookup_sheet["C1"] = "IS_POSTABLE"
    lookup_sheet["D1"] = "BALANCE_TYPE"
    lookup_sheet["E1"] = "SAMPLE_GROUPS"

    for i, c in enumerate(classes, start=2):
        lookup_sheet[f"A{i}"] = c
    for i, cur in enumerate(currencies, start=2):
        lookup_sheet[f"B{i}"] = cur
    for i, p in enumerate(postable_opts, start=2):
        lookup_sheet[f"C{i}"] = p
    for i, b in enumerate(balance_types, start=2):
        lookup_sheet[f"D{i}"] = b
    for i, g in enumerate(sample_groups, start=2):
        lookup_sheet[f"E{i}"] = g

    lookup_sheet.sheet_state = "hidden"

    # Styling helpers
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid") # Teal / Brand
    sample_fill = PatternFill(start_color="F0FDFA", end_color="F0FDFA", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # 2. Accounts Sheet (Primary)
    acct_sheet = wb.create_sheet(title="Accounts")
    acct_headers = [
        "GroupLevel1*",       # Account Class / Nature: Asset, Liability, Equity, Income, Expense
        "AccountGroupName*",  # Account group name
        "AccountNumber",      # Optional code or auto-generated
        "AccountName*",       # Ledger account name
        "Currency*",          # Currency code
        "IsPostable*",        # Yes / No
        "BalanceType",        # DEBIT / CREDIT
        "Description",        # Description / Notes
    ]
    acct_sheet.append(acct_headers)

    # Sample rows for Accounts
    sample_accounts = [
        ["Asset", "Cash and Cash Equivalents", "1001", "Main Petty Cash", "GHS", "Yes", "DEBIT", "Petty cash float"],
        ["Asset", "Bank Accounts", "1002", "GCB Operating Account", "GHS", "Yes", "DEBIT", "Main operating bank account"],
        ["Asset", "Accounts Receivable", "1003", "Trade Debtors Control", "GHS", "Yes", "DEBIT", "Customer receivable ledger"],
        ["Liability", "Accounts Payable", "2001", "Trade Creditors Control", "GHS", "Yes", "CREDIT", "Supplier payable ledger"],
        ["Liability", "Current Liabilities", "2002", "Accrued Utilities", "GHS", "Yes", "CREDIT", "Monthly utility accruals"],
        ["Equity", "Owner's Equity", "3001", "Retained Earnings", "GHS", "Yes", "CREDIT", "Accumulated profit/loss"],
        ["Income", "Operating Revenue", "4001", "Sales Revenue - Goods", "GHS", "Yes", "CREDIT", "Main commercial sales"],
        ["Expense", "Operating Expenses", "5001", "Office Supplies Expense", "GHS", "Yes", "DEBIT", "Stationery and supplies"],
        ["Expense", "Administrative Expenses", "5002", "Rent Expense", "GHS", "Yes", "DEBIT", "Office facility lease"],
    ]
    for row in sample_accounts:
        acct_sheet.append(row)

    # 3. AccountGroups Sheet
    ag_sheet = wb.create_sheet(title="AccountGroups")
    ag_headers = [
        "GroupLevel1*",      # Account Class: Asset, Liability, Equity, Income, Expense
        "GroupName*",        # Group Name
        "GroupCode",         # Group Code (e.g. AST_CURR or empty)
        "ParentGroup",       # Parent group name or empty
        "Description",       # Description
    ]
    ag_sheet.append(ag_headers)

    sample_groups_rows = [
        ["Asset", "Current Assets", "AST_CURR", "", "Short-term operating assets"],
        ["Asset", "Cash and Cash Equivalents", "AST_CASH", "Current Assets", "Liquid funds and bank balances"],
        ["Asset", "Accounts Receivable", "AST_AR", "Current Assets", "Trade receivables and customer debts"],
        ["Liability", "Current Liabilities", "LIA_CURR", "", "Short-term obligations"],
        ["Liability", "Accounts Payable", "LIA_AP", "Current Liabilities", "Trade payables to suppliers"],
        ["Equity", "Owner's Equity", "EQU_OWNER", "", "Capital contributions and reserves"],
        ["Income", "Operating Revenue", "INC_REV", "", "Revenue generated from core operations"],
        ["Expense", "Operating Expenses", "EXP_OPS", "", "Day-to-day operating overheads"],
    ]
    for row in sample_groups_rows:
        ag_sheet.append(row)

    # Style header and sample rows
    for ws in [acct_sheet, ag_sheet]:
        # Header formatting
        for col_idx in range(1, len(ws[1]) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        ws.row_dimensions[1].height = 28

        # Sample rows formatting
        for r in range(2, ws.max_row + 1):
            ws.row_dimensions[r].height = 20
            for col_idx in range(1, len(ws[1]) + 1):
                cell = ws.cell(row=r, column=col_idx)
                cell.fill = sample_fill
                cell.border = thin_border

        # Adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

    # Add Data Validations for Accounts sheet
    # Account Class dropdown (GroupLevel1)
    dv_class = DataValidation(
        type="list",
        formula1="LookupData!$A$2:$A$6",
        allow_blank=False,
        showErrorMessage=True,
        errorTitle="Invalid Account Class",
        error="Please select an Account Class (Asset, Liability, Equity, Income, Expense) from the list."
    )
    acct_sheet.add_data_validation(dv_class)
    dv_class.add("A2:A1000")

    # Currency dropdown
    dv_currency = DataValidation(
        type="list",
        formula1=f"LookupData!$B$2:$B${len(currencies)+1}",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Invalid Currency",
        error="Please select a Currency from the list."
    )
    acct_sheet.add_data_validation(dv_currency)
    dv_currency.add("E2:E1000")

    # IsPostable dropdown
    dv_postable = DataValidation(
        type="list",
        formula1="LookupData!$C$2:$C$3",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Invalid Postable Option",
        error="Please select Yes or No."
    )
    acct_sheet.add_data_validation(dv_postable)
    dv_postable.add("F2:F1000")

    # BalanceType dropdown
    dv_balance = DataValidation(
        type="list",
        formula1="LookupData!$D$2:$D$3",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Invalid Balance Type",
        error="Please select DEBIT or CREDIT."
    )
    acct_sheet.add_data_validation(dv_balance)
    dv_balance.add("G2:G1000")

    # Add Data Validations for AccountGroups sheet
    dv_ag_class = DataValidation(
        type="list",
        formula1="LookupData!$A$2:$A$6",
        allow_blank=False,
        showErrorMessage=True,
        errorTitle="Invalid Account Class",
        error="Please select an Account Class (Asset, Liability, Equity, Income, Expense) from the list."
    )
    ag_sheet.add_data_validation(dv_ag_class)
    dv_ag_class.add("A2:A500")

    # Save to all target paths
    for p in output_paths:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        wb.save(p)
        print(f"Generated template at: {p}")

if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.abspath(__file__))
    targets = [
        os.path.join(base_dir, "account_template.xlsx"),
        os.path.abspath(os.path.join(base_dir, "..", "client", "public", "templates", "account_template.xlsx")),
    ]
    dist_path = os.path.abspath(os.path.join(base_dir, "..", "client", "dist", "templates", "account_template.xlsx"))
    if os.path.exists(os.path.dirname(dist_path)):
        targets.append(dist_path)

    create_account_template(targets)
